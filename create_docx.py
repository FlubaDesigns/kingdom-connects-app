from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# Create document
doc = Document()

# Read markdown
with open('DATABASE-BLING-CHECKLIST.md', 'r') as f:
    lines = f.readlines()

def add_heading(text, level):
    heading = doc.add_heading(text, level=level)
    if level == 1:
        heading.runs[0].font.color.rgb = RGBColor(44, 62, 80)
    elif level == 2:
        heading.runs[0].font.color.rgb = RGBColor(52, 73, 94)

def add_paragraph(text, bold=False, code=False):
    p = doc.add_paragraph()
    if code:
        p.style = 'Intense Quote'
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    if code:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)

def add_code_block(code_lines):
    p = doc.add_paragraph()
    p.style = 'Intense Quote'
    run = p.add_run('\n'.join(code_lines))
    run.font.name = 'Courier New'
    run.font.size = Pt(9)

in_code_block = False
code_block_lines = []

for line in lines:
    line = line.rstrip()
    
    # Code blocks
    if line.startswith('```'):
        if in_code_block:
            add_code_block(code_block_lines)
            code_block_lines = []
            in_code_block = False
        else:
            in_code_block = True
        continue
    
    if in_code_block:
        code_block_lines.append(line)
        continue
    
    # Headers
    if line.startswith('# '):
        add_heading(line[2:], 1)
    elif line.startswith('## '):
        add_heading(line[3:], 2)
    elif line.startswith('### '):
        add_heading(line[4:], 3)
    # Horizontal rule
    elif line == '---':
        doc.add_paragraph('_' * 80)
    # Empty lines
    elif not line.strip():
        continue
    # Regular text
    else:
        # Clean markdown formatting
        text = line
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        
        is_bold = '**' in line or line.startswith('-') or line.startswith('*')
        add_paragraph(text, bold=is_bold)

# Save
doc.save('DATABASE-BLING-CHECKLIST.docx')
print("✅ Created DATABASE-BLING-CHECKLIST.docx")
